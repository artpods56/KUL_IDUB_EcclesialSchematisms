from io import BytesIO
from pathlib import Path
from typing import cast
from uuid import UUID

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from notarius_core.artifacts import ArtifactObject, InMemoryUnitOfWork, JsonObject
from notarius_core.domain.staged_uploads import StagedUpload
from notarius_core.nodes import NodeExecutionContext
from notarius_core.operators.tables import (
    FuzzyMatchTablesNode,
    FuzzyMatchScorer,
    NormalizeTableTextNode,
    TABLES,
    TABLE_DATA,
    Table,
    TableArtifactResolver,
    TableArtifactWriter,
    TableColumn,
    TableFileImportConfig,
    TableFileImportError,
    TableFileImportInput,
    TableFileImportNode,
    TableFileUploadItem,
    TableFuzzyMatchConfig,
    TableFuzzyMatchInput,
    TableManifest,
    TableTextNormalizeConfig,
    TableTextNormalizeInput,
    TableValueType,
    load_table_artifact,
    load_table_manifest,
    load_table_page,
    table_artifact_is_accessible,
)
from notarius_core.plugins import PluginRegistry, PluginRuntimeContext
from notarius_core.ports.storage import SaveFileCommand
from notarius_core.runtime.materialization import MaterializationProvenance
from notarius_core.runtime.persistence import ArtifactWriteContext
from notarius_core.runtime.resolvers import ResolutionError
from notarius_storage import LocalFileObjectStore


TEST_WORKSPACE_ID = UUID("00000000-0000-0000-0000-000000000901")


async def seed_staged_upload(
    unit_of_work: InMemoryUnitOfWork,
    *,
    workspace_id: UUID,
    upload_key: str,
    filename: str,
    byte_size: int,
) -> None:
    async with unit_of_work as entered:
        await entered.staged_uploads.add(
            StagedUpload(
                workspace_id=workspace_id,
                upload_key=upload_key,
                original_filename=filename,
                byte_size=byte_size,
            )
        )
        await entered.commit()


def sample_table() -> Table:
    return Table(
        columns=[
            TableColumn(
                id="column_1",
                title="Name",
                value_type=TableValueType.TEXT,
            ),
            TableColumn(
                id="column_2",
                title="Total",
                value_type=TableValueType.DECIMAL,
            ),
        ],
        rows=[
            {"column_1": "Invoice", "column_2": "19.99"},
            {"column_1": "Receipt", "column_2": None},
        ],
    )


def test_table_model_preserves_duplicate_titles_with_stable_column_ids() -> None:
    table = Table(
        columns=[
            TableColumn(
                id="column_1",
                title="name",
                value_type=TableValueType.TEXT,
            ),
            TableColumn(
                id="column_2",
                title="name",
                value_type=TableValueType.TEXT,
            ),
        ],
        rows=[{"column_1": "first", "column_2": "second"}],
    )

    assert [column.id for column in table.columns] == ["column_1", "column_2"]
    assert [column.title for column in table.columns] == ["name", "name"]


def test_table_model_rejects_non_rectangular_or_mistyped_rows() -> None:
    columns = [
        TableColumn(
            id="count",
            title="Count",
            value_type=TableValueType.INTEGER,
        ),
    ]

    with pytest.raises(ValidationError, match="missing.*count"):
        Table(columns=columns, rows=[{}])
    with pytest.raises(ValidationError, match="unexpected.*other"):
        Table(columns=columns, rows=[{"count": 1, "other": 2}])
    with pytest.raises(ValidationError, match="declares 'integer'.*str"):
        Table(columns=columns, rows=[{"count": "1"}])


@pytest.mark.asyncio
async def test_table_file_import_reads_utf8_csv_with_stable_column_ids(
    tmp_path: Path,
) -> None:
    uploads_dir = tmp_path / "uploads"
    workspace_uploads = uploads_dir / str(TEST_WORKSPACE_ID)
    workspace_uploads.mkdir(parents=True)
    content = (
        "Nr,Miejscowość,Powiat\n1,м. Бѣлыничи,mohylewski\n2,Вендорож,mohylewski\n"
    ).encode()
    upload_path = workspace_uploads / "places.csv"
    upload_path.write_bytes(content)
    uow = InMemoryUnitOfWork()
    await seed_staged_upload(
        uow,
        workspace_id=TEST_WORKSPACE_ID,
        upload_key=upload_path.name,
        filename="places.csv",
        byte_size=len(content),
    )
    node = TableFileImportNode(uploads_dir=uploads_dir, unit_of_work=uow)

    output = await node.run(
        NodeExecutionContext(workspace_id=TEST_WORKSPACE_ID, node_id="table-file"),
        TableFileImportConfig(
            uploads=[
                TableFileUploadItem(
                    upload_key=upload_path.name,
                    filename="places.csv",
                    byte_size=len(content),
                )
            ]
        ),
        TableFileImportInput(),
    )

    assert [column.id for column in output.table.columns] == [
        "column_1",
        "column_2",
        "column_3",
    ]
    assert [column.title for column in output.table.columns] == [
        "Nr",
        "Miejscowość",
        "Powiat",
    ]
    assert output.table.rows == [
        {
            "column_1": "1",
            "column_2": "м. Бѣлыничи",
            "column_3": "mohylewski",
        },
        {
            "column_1": "2",
            "column_2": "Вендорож",
            "column_3": "mohylewski",
        },
    ]


@pytest.mark.asyncio
async def test_table_file_import_selects_xlsx_sheet_and_preserves_scalars(
    tmp_path: Path,
) -> None:
    uploads_dir = tmp_path / "uploads"
    workspace_uploads = uploads_dir / str(TEST_WORKSPACE_ID)
    workspace_uploads.mkdir(parents=True)
    workbook = Workbook()
    ignored = workbook.active
    assert ignored is not None
    ignored.title = "Ignored"
    selected = workbook.create_sheet("Places")
    selected.append(["Name", "Population", "Seat"])
    selected.append(["Belynichi", 10, True])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    content = buffer.getvalue()
    upload_path = workspace_uploads / "places.xlsx"
    upload_path.write_bytes(content)
    uow = InMemoryUnitOfWork()
    await seed_staged_upload(
        uow,
        workspace_id=TEST_WORKSPACE_ID,
        upload_key=upload_path.name,
        filename="places.xlsx",
        byte_size=len(content),
    )
    node = TableFileImportNode(uploads_dir=uploads_dir, unit_of_work=uow)

    output = await node.run(
        NodeExecutionContext(workspace_id=TEST_WORKSPACE_ID, node_id="table-file"),
        TableFileImportConfig(
            uploads=[
                TableFileUploadItem(
                    upload_key=upload_path.name,
                    filename="places.xlsx",
                    byte_size=len(content),
                )
            ],
            sheet_name="Places",
        ),
        TableFileImportInput(),
    )

    assert [column.value_type for column in output.table.columns] == [
        TableValueType.TEXT,
        TableValueType.INTEGER,
        TableValueType.BOOLEAN,
    ]
    assert output.table.rows == [
        {"column_1": "Belynichi", "column_2": 10, "column_3": True}
    ]


@pytest.mark.asyncio
async def test_table_file_import_fails_closed_without_db_row(tmp_path: Path) -> None:
    uploads_dir = tmp_path / "uploads"
    workspace_uploads = uploads_dir / str(TEST_WORKSPACE_ID)
    workspace_uploads.mkdir(parents=True)
    content = b"a,b\n1,2\n"
    upload_path = workspace_uploads / "orphan.csv"
    upload_path.write_bytes(content)
    node = TableFileImportNode(
        uploads_dir=uploads_dir,
        unit_of_work=InMemoryUnitOfWork(),
    )

    with pytest.raises(TableFileImportError, match="was not found in workspace"):
        await node.run(
            NodeExecutionContext(workspace_id=TEST_WORKSPACE_ID, node_id="table-file"),
            TableFileImportConfig(
                uploads=[
                    TableFileUploadItem(
                        upload_key=upload_path.name,
                        filename="orphan.csv",
                        byte_size=len(content),
                    )
                ]
            ),
            TableFileImportInput(),
        )


@pytest.mark.asyncio
async def test_table_file_import_fails_closed_for_foreign_workspace_row(
    tmp_path: Path,
) -> None:
    uploads_dir = tmp_path / "uploads"
    other_workspace = UUID("00000000-0000-0000-0000-000000000902")
    workspace_uploads = uploads_dir / str(TEST_WORKSPACE_ID)
    workspace_uploads.mkdir(parents=True)
    content = b"a,b\n1,2\n"
    upload_path = workspace_uploads / "places.csv"
    upload_path.write_bytes(content)
    uow = InMemoryUnitOfWork()
    await seed_staged_upload(
        uow,
        workspace_id=other_workspace,
        upload_key=upload_path.name,
        filename="places.csv",
        byte_size=len(content),
    )
    node = TableFileImportNode(uploads_dir=uploads_dir, unit_of_work=uow)

    with pytest.raises(TableFileImportError, match="was not found in workspace"):
        await node.run(
            NodeExecutionContext(workspace_id=TEST_WORKSPACE_ID, node_id="table-file"),
            TableFileImportConfig(
                uploads=[
                    TableFileUploadItem(
                        upload_key=upload_path.name,
                        filename="places.csv",
                        byte_size=len(content),
                    )
                ]
            ),
            TableFileImportInput(),
        )


@pytest.mark.asyncio
async def test_text_normalization_adds_transliteration_without_replacing_source() -> (
    None
):
    source = Table(
        columns=[TableColumn(id="name", title="Name", value_type=TableValueType.TEXT)],
        rows=[{"name": "м. Бѣлыничи"}, {"name": None}],
    )

    output = await NormalizeTableTextNode().run(
        NodeExecutionContext(workspace_id=TEST_WORKSPACE_ID, node_id="normalize"),
        TableTextNormalizeConfig(source_column="Name"),
        TableTextNormalizeInput(table=source),
    )

    assert [column.id for column in output.table.columns] == [
        "name",
        "normalized_name",
    ]
    assert output.table.rows == [
        {"name": "м. Бѣлыничи", "normalized_name": "m belynichi"},
        {"name": None, "normalized_name": None},
    ]


@pytest.mark.asyncio
async def test_fuzzy_match_returns_ranked_candidates_and_unmatched_sources() -> None:
    left = Table(
        columns=[
            TableColumn(id="id", title="ID", value_type=TableValueType.INTEGER),
            TableColumn(id="name", title="Name", value_type=TableValueType.TEXT),
            TableColumn(
                id="district",
                title="District",
                value_type=TableValueType.TEXT,
            ),
        ],
        rows=[
            {"id": 1, "name": "belynichi", "district": "mohylewski"},
            {"id": 2, "name": "wendoroz", "district": "mohylewski"},
            {"id": 3, "name": "missing", "district": "mohylewski"},
        ],
    )
    right = Table(
        columns=[
            TableColumn(id="id", title="ID", value_type=TableValueType.INTEGER),
            TableColumn(id="name", title="Name", value_type=TableValueType.TEXT),
            TableColumn(
                id="district",
                title="District",
                value_type=TableValueType.TEXT,
            ),
            TableColumn(
                id="description",
                title="Description",
                value_type=TableValueType.TEXT,
            ),
        ],
        rows=[
            {
                "id": 10,
                "name": "belynichi",
                "district": "mohylewski",
                "description": "SGKP entry",
            },
            {
                "id": 11,
                "name": "belynichi",
                "district": "homelski",
                "description": "Wrong district",
            },
            {
                "id": 12,
                "name": "wendoroz",
                "district": "mohylewski",
                "description": "Historical entry",
            },
        ],
    )

    output = await FuzzyMatchTablesNode().run(
        NodeExecutionContext(workspace_id=TEST_WORKSPACE_ID, node_id="fuzzy-match"),
        TableFuzzyMatchConfig(
            left_text_column="name",
            right_text_column="name",
            left_block_column="district",
            right_block_column="district",
            scorer=FuzzyMatchScorer.RATIO,
            score_threshold=90.0,
            max_candidates=2,
        ),
        TableFuzzyMatchInput(left=left, right=right),
    )

    assert len(output.matches.rows) == 3
    assert output.matches.rows[0]["match_score"] == 100.0
    assert output.matches.rows[0]["match_rank"] == 1
    assert output.matches.rows[0]["right__id"] == 10
    assert output.matches.rows[0]["right__description"] == "SGKP entry"
    assert output.matches.rows[1]["right__id"] == 12
    assert output.matches.rows[2]["left__id"] == 3
    assert output.matches.rows[2]["right_row_index"] is None
    assert output.matches.rows[2]["match_rank"] is None


@pytest.mark.asyncio
async def test_fuzzy_match_uses_alias_columns_and_reports_the_best_pair() -> None:
    left = Table(
        columns=[
            TableColumn(id="name", title="Name", value_type=TableValueType.TEXT),
        ],
        rows=[{"name": "belynichi"}],
    )
    right = Table(
        columns=[
            TableColumn(
                id="historical_name",
                title="Historical name",
                value_type=TableValueType.TEXT,
            ),
            TableColumn(
                id="current_name",
                title="Current name",
                value_type=TableValueType.TEXT,
            ),
        ],
        rows=[
            {
                "historical_name": "beliki",
                "current_name": None,
            },
            {
                "historical_name": "bialenicze",
                "current_name": "byalynichy",
            },
        ],
    )

    output = await FuzzyMatchTablesNode().run(
        NodeExecutionContext(
            workspace_id=TEST_WORKSPACE_ID,
            node_id="fuzzy-match-aliases",
        ),
        TableFuzzyMatchConfig(
            left_text_column="name",
            right_text_column="historical_name",
            right_alias_columns=["current_name"],
            scorer=FuzzyMatchScorer.RATIO,
            score_threshold=50.0,
            max_candidates=2,
        ),
        TableFuzzyMatchInput(left=left, right=right),
    )

    assert output.matches.rows[0]["right__historical_name"] == "bialenicze"
    assert output.matches.rows[0]["match_right_column"] == "current_name"
    assert output.matches.rows[1]["right__historical_name"] == "beliki"
    assert output.matches.rows[1]["match_right_column"] == "historical_name"


def test_table_plugin_registers_chunked_persistence(tmp_path: Path) -> None:
    registry = PluginRegistry()
    registry.install(TABLES)
    registry.freeze()
    context = PluginRuntimeContext(
        workspace=tmp_path,
        uploads_dir=tmp_path / "uploads",
        storage=LocalFileObjectStore(tmp_path / "artifacts"),
        uow=InMemoryUnitOfWork(),
        bucket="artifacts",
    )

    assert TABLES.slug == "builtin.table"
    assert TABLES.title == "Table"
    assert [artifact.key for artifact in registry.artifact_types] == [TABLE_DATA.key]
    assert TABLE_DATA.payload_schema == Table.model_json_schema()
    resolver = registry.build_resolvers(context)[0]
    writer = registry.build_writers(context)[0]
    assert isinstance(resolver, TableArtifactResolver)
    assert resolver.target is Table
    assert isinstance(writer, TableArtifactWriter)
    assert writer.artifact_type == TABLE_DATA.key
    assert sample_table().rows[0]["column_2"] == "19.99"
    assert {node.key for node in registry.nodes} >= {
        ("table.file.import", 1),
        ("table.text.normalize", 1),
        ("table.fuzzy_match", 1),
    }


@pytest.mark.asyncio
async def test_chunked_table_round_trip_and_cross_chunk_page(tmp_path: Path) -> None:
    unit_of_work = InMemoryUnitOfWork()
    storage = LocalFileObjectStore(tmp_path / "artifacts")
    writer = TableArtifactWriter(
        storage=storage,
        uow=unit_of_work,
        bucket="artifacts",
        storage_backend="local",
    )
    table = Table(
        columns=[
            TableColumn(id="row", title="Row", value_type=TableValueType.INTEGER),
            TableColumn(id="value", title="Value", value_type=TableValueType.TEXT),
        ],
        rows=[{"row": index, "value": f"value-{index}"} for index in range(205)],
    )

    ref = await writer.write(
        table,
        ArtifactWriteContext(
            node_context=NodeExecutionContext(
                workspace_id=TEST_WORKSPACE_ID,
                node_id="table",
            ),
            provenance=MaterializationProvenance(refs_by_input={}),
        ),
    )
    async with unit_of_work as uow:
            artifact = await uow.artifacts.get(TEST_WORKSPACE_ID, ref.artifact_id)
            assert artifact is not None
            assert artifact.workspace_id == TEST_WORKSPACE_ID
            assert artifact.object_key is not None
            assert artifact.object_key.startswith(
                f"workspaces/{TEST_WORKSPACE_ID}/table.data/v1/manifests/"
            )
    assert artifact is not None
    assert artifact.inline_payload is None
    assert artifact.metadata["row_count"] == 205
    assert artifact.metadata["column_count"] == 2
    assert artifact.metadata["chunk_count"] == 3

    page = await load_table_page(artifact, storage, offset=95, limit=10)
    resolver = TableArtifactResolver(uow=unit_of_work, storage=storage)

    assert page.offset == 95
    assert page.total_rows == 205
    assert [row["row"] for row in page.rows] == list(range(95, 105))
    past_end = await load_table_page(artifact, storage, offset=999, limit=10)
    assert past_end.offset == 205
    assert past_end.rows == []
    assert await resolver.resolve(ref, TEST_WORKSPACE_ID) == table


@pytest.mark.asyncio
async def test_empty_and_legacy_inline_tables_remain_resolvable(tmp_path: Path) -> None:
    unit_of_work = InMemoryUnitOfWork()
    storage = LocalFileObjectStore(tmp_path / "artifacts")
    writer = TableArtifactWriter(
        storage=storage,
        uow=unit_of_work,
        bucket="artifacts",
        storage_backend="local",
    )
    empty = Table(
        columns=[
            TableColumn(id="value", title="Value", value_type=TableValueType.TEXT)
        ],
        rows=[],
    )
    empty_ref = await writer.write(
        empty,
        ArtifactWriteContext(
            node_context=NodeExecutionContext(
                workspace_id=TEST_WORKSPACE_ID,
                node_id="empty-table",
            ),
            provenance=MaterializationProvenance(refs_by_input={}),
        ),
    )
    legacy = sample_table()
    legacy_artifact = ArtifactObject(
        workspace_id=TEST_WORKSPACE_ID,
        artifact_type=TABLE_DATA.key.id,
        schema_version=TABLE_DATA.key.schema_version,
        content_type="application/json",
        storage_backend="inline",
        inline_payload=cast(JsonObject, legacy.model_dump(mode="json")),
    )
    async with unit_of_work as uow:
        await uow.artifacts.add(legacy_artifact)
        await uow.commit()

    resolver = TableArtifactResolver(uow=unit_of_work, storage=storage)
    legacy_page = await load_table_page(
        legacy_artifact,
        storage,
        offset=1,
        limit=10,
    )

    assert await resolver.resolve(empty_ref, TEST_WORKSPACE_ID) == empty
    assert await resolver.resolve(legacy_artifact.ref(), TEST_WORKSPACE_ID) == legacy
    assert legacy_page.total_rows == 2
    assert legacy_page.rows == legacy.rows[1:]


@pytest.mark.asyncio
async def test_corrupt_table_chunk_reports_artifact_and_offset(tmp_path: Path) -> None:
    unit_of_work = InMemoryUnitOfWork()
    storage = LocalFileObjectStore(tmp_path / "artifacts")
    writer = TableArtifactWriter(
        storage=storage,
        uow=unit_of_work,
        bucket="artifacts",
        storage_backend="local",
    )
    ref = await writer.write(
        sample_table(),
        ArtifactWriteContext(
            node_context=NodeExecutionContext(
                workspace_id=TEST_WORKSPACE_ID,
                node_id="table",
            ),
            provenance=MaterializationProvenance(refs_by_input={}),
        ),
    )
    async with unit_of_work as uow:
            artifact = await uow.artifacts.get(TEST_WORKSPACE_ID, ref.artifact_id)
    assert artifact is not None
    assert artifact.bucket is not None
    manifest = await load_table_manifest(artifact, storage)
    chunk = manifest.chunks[0]
    await storage.save(
        SaveFileCommand(
            bucket=artifact.bucket,
            path=chunk.object_key,
            stream=BytesIO(b"{}"),
            content_type="application/json",
            metadata={},
            allow_overwrite=True,
        )
    )

    with pytest.raises(
        ResolutionError,
        match=rf"offset 0.*artifact {artifact.id}",
    ):
        await load_table_page(artifact, storage, offset=0, limit=1)


@pytest.mark.asyncio
async def test_table_accessibility_requires_every_chunk(tmp_path: Path) -> None:
    unit_of_work = InMemoryUnitOfWork()
    storage = LocalFileObjectStore(tmp_path / "artifacts")
    writer = TableArtifactWriter(
        storage=storage,
        uow=unit_of_work,
        bucket="artifacts",
        storage_backend="local",
    )
    table = Table(
        columns=[TableColumn(id="row", title="Row", value_type=TableValueType.INTEGER)],
        rows=[{"row": index} for index in range(205)],
    )
    ref = await writer.write(
        table,
        ArtifactWriteContext(
            node_context=NodeExecutionContext(
                workspace_id=TEST_WORKSPACE_ID,
                node_id="table",
            ),
            provenance=MaterializationProvenance(refs_by_input={}),
        ),
    )
    async with unit_of_work as uow:
            artifact = await uow.artifacts.get(TEST_WORKSPACE_ID, ref.artifact_id)
    assert artifact is not None
    assert artifact.bucket is not None
    manifest = await load_table_manifest(artifact, storage)

    assert await table_artifact_is_accessible(artifact, storage)
    await storage.delete(artifact.bucket, manifest.chunks[1].object_key)
    assert not await table_artifact_is_accessible(artifact, storage)


@pytest.mark.asyncio
async def test_large_rows_are_split_by_chunk_byte_budget(tmp_path: Path) -> None:
    unit_of_work = InMemoryUnitOfWork()
    storage = LocalFileObjectStore(tmp_path / "artifacts")
    writer = TableArtifactWriter(
        storage=storage,
        uow=unit_of_work,
        bucket="artifacts",
        storage_backend="local",
    )
    table = Table(
        columns=[
            TableColumn(
                id="geometry",
                title="Geometry",
                value_type=TableValueType.TEXT,
            )
        ],
        rows=[{"geometry": f"{index}-" + "x" * 600_000} for index in range(4)],
    )
    ref = await writer.write(
        table,
        ArtifactWriteContext(
            node_context=NodeExecutionContext(
                workspace_id=TEST_WORKSPACE_ID,
                node_id="large-table",
            ),
            provenance=MaterializationProvenance(refs_by_input={}),
        ),
    )
    async with unit_of_work as uow:
            artifact = await uow.artifacts.get(TEST_WORKSPACE_ID, ref.artifact_id)
    assert artifact is not None
    manifest = await load_table_manifest(artifact, storage)
    page = await load_table_page(artifact, storage, offset=1, limit=2)

    assert [chunk.row_count for chunk in manifest.chunks] == [1, 1, 1, 1]
    assert page.rows == table.rows[1:3]


@pytest.mark.asyncio
async def test_table_manifest_and_logical_content_are_authenticated(
    tmp_path: Path,
) -> None:
    unit_of_work = InMemoryUnitOfWork()
    storage = LocalFileObjectStore(tmp_path / "artifacts")
    writer = TableArtifactWriter(
        storage=storage,
        uow=unit_of_work,
        bucket="artifacts",
        storage_backend="local",
    )
    ref = await writer.write(
        sample_table(),
        ArtifactWriteContext(
            node_context=NodeExecutionContext(
                workspace_id=TEST_WORKSPACE_ID,
                node_id="table",
            ),
            provenance=MaterializationProvenance(refs_by_input={}),
        ),
    )
    async with unit_of_work as uow:
            artifact = await uow.artifacts.get(TEST_WORKSPACE_ID, ref.artifact_id)
    assert artifact is not None
    assert artifact.bucket is not None
    assert artifact.object_key is not None
    manifest = await load_table_manifest(artifact, storage)
    replacement_manifest = TableManifest(
        columns=[
            column.model_copy(update={"title": f"Altered {column.title}"})
            for column in manifest.columns
        ],
        row_count=manifest.row_count,
        chunks=manifest.chunks,
    )
    await storage.save(
        SaveFileCommand(
            bucket=artifact.bucket,
            path=artifact.object_key,
            stream=BytesIO(replacement_manifest.model_dump_json().encode("utf-8")),
            content_type="application/json",
            metadata={},
            allow_overwrite=True,
        )
    )

    with pytest.raises(ResolutionError, match=rf"manifest.*artifact {artifact.id}"):
        await load_table_manifest(artifact, storage)

    await storage.save(
        SaveFileCommand(
            bucket=artifact.bucket,
            path=artifact.object_key,
            stream=BytesIO(manifest.model_dump_json().encode("utf-8")),
            content_type="application/json",
            metadata={},
            allow_overwrite=True,
        )
    )
    artifact.sha256 = "0" * 64
    with pytest.raises(
        ResolutionError,
        match=rf"artifact {artifact.id}.*SHA-256",
    ):
        await load_table_artifact(artifact, storage)
